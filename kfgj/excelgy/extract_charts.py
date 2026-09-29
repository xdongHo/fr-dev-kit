#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
extract_charts.py — 流程 F 第⑤步：从 Excel 抽取图表几何/类型/系列 → 布局 JSON

输出对齐 ckzl/hbfa 的机器可读方案（containers: x/y/w/h + components），
供后续映射到 cpt 单元格内嵌图表（<O t="CC">）或 frm 容器（网页框/报表块）。

坐标换算：
  - 锚点 EMU → px（1px = 9525 EMU）
  - 列宽/行高优先用工作表实际值（openpyxl width≈7px/单位, height≈px×0.75），
    未设置时按 canvas 约定（列 24px / 行 20px，即 frm 画布式 px÷24/÷20 的逆运算）

用法：
    python kfgj/excelgy/extract_charts.py <book.xlsx> [-o layout.json] [--canvas 1280x768]
"""
import argparse
import io
import json
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from openpyxl import load_workbook

EMU_PER_PX = 9525
TYPE_MAP = {
    "BarChart": "bar", "LineChart": "line", "PieChart": "pie",
    "AreaChart": "area", "ScatterChart": "scatter", "DoughnutChart": "doughnut",
    "RadarChart": "radar", "StockChart": "stock", "SurfaceChart": "surface",
}
FR_TYPE = {"bar": "柱形图", "line": "折线图", "pie": "饼图", "area": "面积图",
           "scatter": "散点图", "doughnut": "圆环图", "radar": "雷达图", "stock": "股价图"}


def col_px(ws, ci):
    """第 ci 列(1基)的像素宽：实际 width 未设时回退 24px（canvas 约定）"""
    letter = ws.cell(row=1, column=ci).column_letter
    cd = ws.column_dimensions.get(letter)
    if cd is not None and cd.width:
        return cd.width * 7
    return 24.0


def row_px(ws, ri):
    rd = ws.row_dimensions.get(ri)
    if rd is not None and rd.height:
        return rd.height / 0.75
    return 20.0


def anchor_to_px(ws, anc):
    """TwoCellAnchor(_from/to) 与 OneCellAnchor(_from/ext) → (x, y, w, h) px"""
    f = anc._from
    x = sum(col_px(ws, c + 1) for c in range(f.col)) + f.colOff / EMU_PER_PX
    y = sum(row_px(ws, r + 1) for r in range(f.row)) + f.rowOff / EMU_PER_PX
    if hasattr(anc, "to") and anc.to is not None:
        t = anc.to
        x2 = sum(col_px(ws, c + 1) for c in range(t.col)) + t.colOff / EMU_PER_PX
        y2 = sum(row_px(ws, r + 1) for r in range(t.row)) + t.rowOff / EMU_PER_PX
        w, h = x2 - x, y2 - y
    else:  # OneCellAnchor: ext(cx, cy) 为 EMU 尺寸
        w, h = anc.ext.cx / EMU_PER_PX, anc.ext.cy / EMU_PER_PX
    return round(x), round(y), round(w), round(h)


def chart_title(ch):
    try:
        for p in ch.title.tx.rich.p:
            for r in (p.r or []):
                if r.t:
                    return r.t
    except Exception:
        pass
    return None


def series_info(ch):
    out = []
    for s in ch.series:
        item = {}
        try:
            if s.tx is not None and s.tx.strRef is not None:
                item["name"] = s.tx.strRef.f
        except Exception:
            pass
        try:
            cat = s.cat.strRef or s.cat.numRef if s.cat is not None else None
            if cat is not None:
                item["category"] = cat.f
        except Exception:
            pass
        try:
            if s.val is not None and s.val.numRef is not None:
                item["value"] = s.val.numRef.f
        except Exception:
            pass
        out.append(item)
    return out


def extract(path, canvas):
    wb = load_workbook(path)
    doc = {"source": path, "canvas": {"width": canvas[0], "height": canvas[1]},
           "sheets": []}
    for ws in wb.worksheets:
        charts = []
        for ch in getattr(ws, "_charts", []):
            ctype = TYPE_MAP.get(type(ch).__name__, type(ch).__name__)
            x, y, w, h = anchor_to_px(ws, ch.anchor)
            entry = {
                "name": chart_title(ch) or f"{ctype}_{len(charts) + 1}",
                "excel_type": ctype,
                "fr_type": FR_TYPE.get(ctype, ctype),
                "x": x, "y": y, "w": w, "h": h,
                "grouping": getattr(ch, "grouping", None),
                "series": series_info(ch),
            }
            charts.append(entry)
        if charts:
            doc["sheets"].append({"name": ws.title, "charts": charts})
    # 汇总成 hbfa 风格 containers（百分比栅格，对齐 hbfa JSON）
    doc["containers"] = [
        {"title": c["name"], "type": "chart", "components": c["fr_type"],
         "x": round(c["x"] * 100 / canvas[0]), "y": round(c["y"] * 100 / canvas[1]),
         "w": round(c["w"] * 100 / canvas[0]), "h": round(c["h"] * 100 / canvas[1]),
         "px": {"x": c["x"], "y": c["y"], "w": c["w"], "h": c["h"]},
         "series": c["series"]}
        for sh in doc["sheets"] for c in sh["charts"]
    ]
    return doc


def main():
    ap = argparse.ArgumentParser(description="Excel 图表 → hbfa 风格布局 JSON")
    ap.add_argument("xlsx")
    ap.add_argument("-o", "--out")
    ap.add_argument("--canvas", default="1280x768",
                    help="画布尺寸，默认 1280x768（frm 骨架库另用 1280×720）")
    a = ap.parse_args()
    w, h = (int(v) for v in a.canvas.lower().split("x"))
    doc = extract(a.xlsx, (w, h))
    text = json.dumps(doc, ensure_ascii=False, indent=1)
    if a.out:
        with open(a.out, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"[extract_charts] ✅ {len(doc['containers'])} 个图表 → {a.out}")
    print(text)


if __name__ == "__main__":
    main()
