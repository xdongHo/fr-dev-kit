#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gen_skeleton.py — 流程 F 第①步：示例数据 + 业务需求契约 → Excel 表样骨架 xlsx

遵循 ckzl/lc/00_需求分析流水线.md §2 表样规范：
  detail 模式（cpt 网格式）：
    - 顶部 5 行 = 参数面板示例（`[参数] 控件名` + 示例值，导入设计器后可删）
    - 表头行 + ds.field 绑定行（蓝字浅蓝底）+ 示例数据行（静态）+ 合计公式行
    - 🟠 [图表 类名] 占位 = 合并单元格
  canvas 模式（frm 画布式，默认 1280×768）：
    - Excel 当画布：坐标 px÷24(列)/÷20(行)
    - 组件 = 合并单元格区域；橙=[图表] 绿=[报表块] 灰=[控件]

继承 GLM-in-Excel 逆向安全协议（~/.fr_dev_kit/experience/2026-09-29 笔记）：
  - 公式风险预检：契约公式 high 级（除零/INDIRECT/OFFSET）直接拒绝
  - 生成后回读断言（写完必验证）

用法：
    python kfgj/excelgy/gen_skeleton.py kfgj/excelgy/example_contract.json -o scwj/mb/骨架.xlsx
"""
import argparse
import io
import json
import os
import re
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── 表样规范色（00_需求分析流水线 §2）──
C_TITLE = "1F4E79"      # 深蓝标题底
C_HEADER = "2E75B6"     # 表头蓝
C_DS_FIELD = "DDEBF7"   # ds.field 绑定格浅蓝底
C_CHART = "FCE4D6"      # 橙 = 图表占位
C_BLOCK = "E2EFDA"      # 绿 = 报表块
C_CTRL = "D9D9D9"       # 灰 = 控件
C_SUM = "D9E2F1"        # 合计行浅蓝

HIGH_RISK = [(r"/\s*0(?![.\d])", "除以字面量 0"),
             (r"\bINDIRECT\s*\(", "INDIRECT 易失引用"),
             (r"\bOFFSET\s*\(", "OFFSET 易失引用")]


def formula_risk(f):
    why = [w for pat, w in HIGH_RISK if re.search(pat, f, re.I)]
    return ("high", why) if why else ("low", [])


def check_contract_formulas(contract):
    """公式风险预检：summary.formula / columns[].formula 中的 Excel 公式"""
    bad = []
    s = contract.get("summary") or {}
    for f in [s.get("formula")] + [c.get("formula") for c in contract.get("columns", [])]:
        if isinstance(f, str) and f.startswith("="):
            lv, why = formula_risk(f)
            if lv == "high":
                bad.append({"formula": f, "reasons": why,
                            "suggestion": "改为 Python 侧计算写入静态值，或改用低风险公式"})
    return bad


def _sty(cell, *, bold=None, size=None, color=None, fill=None, align=None,
         border=False, numfmt=None):
    if bold is not None or size or color:
        cell.font = Font(bold=bold if bold is not None else cell.font.bold,
                         size=size or cell.font.size, color=color or cell.font.color)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if align:
        cell.alignment = Alignment(horizontal=align, vertical="center")
    if border:
        s = Side(style="thin", color="000000")
        cell.border = Border(left=s, right=s, top=s, bottom=s)
    if numfmt:
        cell.number_format = numfmt


def _merge(ws, r1, c1, r2, c2):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def _style_region(ws, r1, c1, r2, c2, **kw):
    """区域内逐格上样式（实体化每个格）。

    铁律 ①：合并/占位区域必须先 merge 再上样式——openpyxl 的 merge_cells()
    会清掉覆盖格上已设的样式并导致这些格不落盘；FR 的 Excel2007ReportImporter
    在 traverseByRows→createStyle 处对"合并区右下角格缺失"未判空（endRow.getCell
    无 null 保护，fine-report-engine 反编译实证 2026-09-29），覆盖格缺失即 NPE。
    铁律 ②：先 merge 后对 MergedCell 上样式可正常保留并落盘（PoC 实证）。
    """
    for ri in range(r1, r2 + 1):
        for ci in range(c1, c2 + 1):
            _sty(ws.cell(row=ri, column=ci), **kw)


# ─────────────────────── detail 模式（cpt 网格式） ───────────────────────
def build_detail(ws, ct):
    cols = ct["columns"]
    ncol = len(cols)
    ds = ct.get("ds_name", "ds1")

    # R1 标题（先 merge 后样式，见 _style_region 铁律）
    _merge(ws, 1, 1, 1, ncol)
    ws.cell(row=1, column=1, value=ct["report"]["name"])
    _style_region(ws, 1, 1, 1, ncol, bold=True, size=14, color="FFFFFF",
                  fill=C_TITLE, align="center")

    # R2/R3 参数面板示例（顶部 5 行区，R4-R5 留空，导入设计器后可删）
    params = ct.get("params", [])
    for i, p in enumerate(params, start=1):
        lc = ws.cell(row=2, column=i, value=f"[参数] {p.get('label', p['name'])}({p['name']})")
        _sty(lc, fill="FFF2CC", border=True)
        ec = ws.cell(row=3, column=i, value=p.get("example", ""))
        _sty(ec, fill="FFF2CC", border=True, align="center")

    # R6 表头
    for j, c in enumerate(cols, start=1):
        hc = ws.cell(row=6, column=j, value=c["title"])
        _sty(hc, bold=True, color="FFFFFF", fill=C_HEADER, align="center", border=True)

    # R7 ds.field 绑定行（蓝字浅蓝底）
    for j, c in enumerate(cols, start=1):
        bind = c.get("bind", "ds")
        val = f"{ds}.{c['field']}" if bind == "ds" else c.get("formula", f"{ds}.{c['field']}")
        bc = ws.cell(row=7, column=j, value=val)
        _sty(bc, color="0070C0", fill=C_DS_FIELD, border=True)

    # R8+ 示例数据行（静态，F4 步骤替换/删除）
    r = 8
    for row in ct.get("sample_rows", []):
        for j, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=j, value=v)
            _sty(cell, border=True,
                 numfmt=cols[j - 1].get("numfmt") if isinstance(v, (int, float)) else None)
        r += 1

    # 合计行
    s = ct.get("summary")
    if s:
        ws.cell(row=r, column=1, value=s.get("label", "合计"))
        _sty(ws.cell(row=r, column=1), bold=True, fill=C_SUM, align="right")
        _merge(ws, r, 1, r, s.get("col", ncol) - 1)
        for cc in range(1, s.get("col", ncol)):
            _sty(ws.cell(row=r, column=cc), fill=C_SUM)
        fc = ws.cell(row=r, column=s.get("col", ncol), value=s["formula"])
        _sty(fc, bold=True, fill=C_SUM, numfmt="#,##0.00", align="right", border=True)

    # 图表占位（🟠 合并格，先 merge 后样式）
    for ch in ct.get("charts", []):
        r2 = ch["row"] + ch.get("rs", 8) - 1
        c2 = ch["col"] + ch.get("cs", 4) - 1
        _merge(ws, ch["row"], ch["col"], r2, c2)
        ws.cell(row=ch["row"], column=ch["col"],
                value=f"[图表 {ch.get('fr_type', '柱形图')}] {ch['name']}")
        _style_region(ws, ch["row"], ch["col"], r2, c2,
                      fill=C_CHART, align="center", border=True)

    for j, c in enumerate(cols, start=1):
        ws.column_dimensions[chr(64 + j) if j <= 26 else "A"].width = c.get("width", 14)
    return ws


# ─────────────────────── canvas 模式（frm 画布式） ───────────────────────
PX_PER_COL = 24   # 表样规范：px÷24(列)
PX_PER_ROW = 20   # 表样规范：px÷20(行)


def build_canvas(ws, ct):
    cv = ct.get("canvas", {})
    W, H = cv.get("width", 1280), cv.get("height", 768)  # 用户约定 1280×768；frm 骨架库另有 1280×720
    ncol, nrow = -(-W // PX_PER_COL), -(-H // PX_PER_ROW)

    for ci in range(1, ncol + 1):
        letter = ws.cell(row=1, column=ci).column_letter
        ws.column_dimensions[letter].width = PX_PER_COL / 7
    for ri in range(1, nrow + 1):
        ws.row_dimensions[ri].height = PX_PER_ROW * 0.75

    fills = {"chart": C_CHART, "block": C_BLOCK, "ctrl": C_CTRL,
             "title": C_TITLE}
    labels = {"chart": "图表", "block": "报表块", "ctrl": "控件", "title": ""}

    for comp in ct.get("charts", []) + ct.get("blocks", []) + ct.get("ctrls", []):
        kind = comp.get("kind", "chart")
        x, y, w, h = comp["x"], comp["y"], comp["w"], comp["h"]
        c1, r1 = x // PX_PER_COL + 1, y // PX_PER_ROW + 1
        c2, r2 = (x + w - 1) // PX_PER_COL + 1, (y + h - 1) // PX_PER_ROW + 1
        text = (f"[{labels[kind]} {comp.get('fr_type', comp.get('widget', ''))}] {comp['name']}"
                if kind != "title" else comp["name"])
        cell = ws.cell(row=r1, column=c1, value=text)
        _sty(cell, fill=fills[kind],
             color="FFFFFF" if kind == "title" else None,
             bold=kind in ("title",), align="center" if kind != "chart" else "left")
        if (r2, c2) != (r1, c1):
            _merge(ws, r1, c1, r2, c2)
    ws.sheet_view.showGridLines = False
    return ws


# ─────────────────────── 主流程 ───────────────────────
def generate(ct, out_path):
    bad = check_contract_formulas(ct)
    if bad:
        print("[gen_skeleton] ❌ 公式风险预检 high 级拦截：")
        for b in bad:
            print(f"   {b['formula']}  原因: {'; '.join(b['reasons'])}  建议: {b['suggestion']}")
        return 1

    wb = Workbook()
    ws = wb.active
    ws.title = ct["report"].get("sheet", "表样")
    mode = ct["report"].get("type", "detail")
    (build_detail if mode == "detail" else build_canvas)(ws, ct)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)

    # 写后回读断言（安全协议：写完必验证）
    vwb = load_workbook(out_path)
    vws = vwb.active
    ok_title = (vws.cell(row=1, column=1).value == ct["report"]["name"]) if mode == "detail" \
        else (vws.max_column >= -(-ct.get("canvas", {}).get("width", 1280) // PX_PER_COL))
    n_sample = len(ct.get("sample_rows", []))
    ok_rows = all(vws.cell(row=8 + i, column=1).value == ct["sample_rows"][i][0]
                  for i in range(n_sample)) if mode == "detail" else True
    print(f"[gen_skeleton] mode={mode} 校验: title={ok_title} sample_rows={ok_rows}")
    print(f"[gen_skeleton] ✅ 生成 {out_path}"
          if ok_title and ok_rows else "[gen_skeleton] ❌ 回读校验失败")
    return 0 if ok_title and ok_rows else 2


def main():
    ap = argparse.ArgumentParser(description="示例数据契约 → Excel 表样骨架")
    ap.add_argument("contract", help="契约 JSON 文件（见 example_contract.json）")
    ap.add_argument("-o", "--out", required=True, help="输出 xlsx 路径")
    a = ap.parse_args()
    with open(a.contract, encoding="utf-8") as f:
        ct = json.load(f)
    return generate(ct, a.out)


if __name__ == "__main__":
    sys.exit(main())
